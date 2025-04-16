# --- Environment Setup ---
from dotenv import load_dotenv
import os

# Load environment variables from .env file
load_dotenv()

# Get API keys
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
LANGCHAIN_API_KEY = os.getenv("LANGCHAIN_API_KEY")

if not GROQ_API_KEY and not OPENAI_API_KEY and not LANGCHAIN_API_KEY:
    raise ValueError("Please set at least one of GROQ_API_KEY, OPENAI_API_KEY, or LANGCHAIN_API_KEY in your .env file.")

# --- Imports ---
from langchain_groq import ChatGroq
from langchain.chat_models import ChatOpenAI
from langchain_community.chat_models import ChatAnthropic
from langchain_core.prompts import ChatPromptTemplate
from langchain.schema import StrOutputParser
from langchain.schema.runnable import Runnable
from langchain.schema.runnable.config import RunnableConfig
from langchain.document_loaders import PyPDFLoader, TextLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
import chainlit as cl
import traceback
import pandas as pd  # File processing
from transformers import pipeline  # Text generation for auto-completion
import speech_recognition as sr
from docx import Document

# Load the text generation model for auto-completion
text_generator = pipeline("text-generation", model="gpt2")

def generate_auto_completion(prompt: str) -> str:
    try:
        completion = text_generator(prompt, max_length=50, num_return_sequences=1, no_repeat_ngram_size=2)[0]['generated_text']
        return completion[len(prompt):].strip()
    except Exception as e:
        return f"Error generating auto-completion: {e}"

# Load the emotion detection model
emotion_detector = pipeline("text-classification", model="j-hartmann/emotion-english-distilroberta-base", return_all_scores=True)

def analyse_emotion(text: str) -> tuple:
    try:
        emotions = emotion_detector(text)[0]
        dominant_emotion = max(emotions, key=lambda x: x['score'])
        return dominant_emotion['label'], dominant_emotion['score']
    except Exception as e:
        return "neutral", 0.0

async def process_csv(file_path: str) -> str:
    try:
        df = pd.read_csv(file_path)
        return df.to_string(index=False)
    except Exception as e:
        return f"Error processing CSV file: {e}"

async def process_excel(file_path: str) -> str:
    try:
        df = pd.read_excel(file_path)
        return df.to_string(index=False)
    except Exception as e:
        return f"Error processing Excel file: {e}"

async def process_word(file_path: str) -> str:
    try:
        doc = Document(file_path)
        if not doc.paragraphs:
            return "The Word document is empty."
        content = "\n".join([paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()])
        return content if content else "The Word document contains no readable text."
    except Exception as e:
        return f"Error processing Word file: {e}"

async def process_file(file_path: str) -> str:
    file_extension = os.path.splitext(file_path)[1].lower()

    if file_extension == ".pdf":
        loader = PyPDFLoader(file_path)
        pages = loader.load()
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        docs = text_splitter.split_documents(pages)
        return "\n\n".join([doc.page_content for doc in docs])
    elif file_extension == ".csv":
        return await process_csv(file_path)
    elif file_extension in [".xls", ".xlsx"]:
        return await process_excel(file_path)
    elif file_extension in [".doc", ".docx"]:
        return await process_word(file_path)
    else:
        return "Unsupported file type. Please upload a PDF, CSV, Excel, or Word (.docx) file."

async def process_voice_input():
    r = sr.Recognizer()
    with sr.Microphone() as source:
        try:
            await cl.Message(content="Listening...").send()
            audio = r.listen(source)
            user_query = r.recognize_google(audio)
            await cl.Message(content=f"You said: {user_query}").send()
            await on_message(cl.Message(content=user_query))
        except sr.UnknownValueError:
            await cl.Message(content="Sorry, I could not understand your speech. Please try again.").send()
        except sr.RequestError:
            await cl.Message(content="There was an issue with the speech recognition service. Please try again later.").send()

@cl.on_chat_start
async def on_chat_start():
    elements = [cl.Image(name="image1", display="inline", path="chat.png")]
    await cl.Message(content="Hello! I am Chatbot. Ask me anything or upload a file.", elements=elements).send()

    backend = os.getenv("LANGCHAIN_BACKEND", "groq").lower()

    if backend == "groq":
        model = ChatGroq(temperature=0, model_name="llama3-70b-8192", groq_api_key=GROQ_API_KEY)
    elif backend == "openai":
        model = ChatOpenAI(temperature=0, model_name="gpt-4", openai_api_key=OPENAI_API_KEY)
    elif backend == "langchain":
        model = ChatAnthropic(temperature=0, model_name="claude-3-opus-20240229", api_key=LANGCHAIN_API_KEY)
    else:
        raise ValueError("Unsupported backend selected.")

    prompt = ChatPromptTemplate.from_messages([
        ("system", "You're a knowledgeable Machine Learning Engineer."),
        ("human", "{question}")
    ])
    cl.user_session.set("runnable", prompt | model | StrOutputParser())

@cl.on_message
async def on_message(message: cl.Message):
    runnable = cl.user_session.get("runnable")
    msg = cl.Message(content="")
    user_query = message.content
    file_text = ""

    if user_query.lower() == "record voice":
        await process_voice_input()
        return

    if message.elements:
        for element in message.elements:
            if isinstance(element, cl.File):
                file_text = await process_file(element.path)
                if "Unsupported file type" in file_text:
                    await cl.Message(content=file_text).send()
                    return

    combined_input = f"User Query: {user_query}\n\nFile Content:\n{file_text}" if file_text else user_query

    if user_query:
        auto_completion = generate_auto_completion(user_query)
        await cl.Message(content=f"Auto-completion: {auto_completion}").send()

    if user_query:
        emotion, emotion_score = analyse_emotion(user_query)
        emotion_icons = {
            "joy": "😊", "anger": "😠", "sadness": "😢",
            "fear": "😨", "surprise": "😲", "love": "❤️",
            "neutral": "😐"
        }
        icon = emotion_icons.get(emotion, "")
        await cl.Message(content=f"Emotion: {emotion} {icon} (Score: {emotion_score:.2f})").send()

    try:
        async for chunk in runnable.astream(
            {"question": combined_input},
            config=RunnableConfig(callbacks=[cl.LangchainCallbackHandler()]),
        ):
            await msg.stream_token(chunk)
    except Exception as e:
        msg.content = f"An error occurred: {traceback.format_exc()}"
        await msg.send()
    else:
        await msg.send()
